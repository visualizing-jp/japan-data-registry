#!/usr/bin/env python3
"""登記統計年報の Excel から public/data の cube を書く。"""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
ESTAT = ROOT / "data/raw/estat"
OUT = ROOT / "public/data"

ERA_FROM = 2010
ERA_TO = 2024
TYPE_FROM = 2015

METRICS = [
    ("total", "登記の総数", "規模"),
    ("real_estate", "不動産", "規模"),
    ("land", "土地", "規模"),
    ("building", "建物", "規模"),
    ("sale", "売買（土地＋建物）", "名義"),
    ("inherit", "相続（土地＋建物）", "名義"),
    ("gift", "贈与（土地＋建物）", "名義"),
    ("mortgage", "抵当権の設定（土地＋建物）", "名義"),
]

TYPE_CODES = [
    ("sale", "売買", "売買による所有権の移転"),
    ("inherit", "相続", None),
    ("gift", "贈与", "遺贈又は贈与による"),
    ("mortgage", "抵当権の設定", "抵当権の設定"),
    ("root", "根抵当権の設定", "根抵当権の設定"),
    ("preserve", "所有権の保存", "所有権の保存"),
    ("display", "表示に関する登記", "土地の表示に関する登記"),
]

COMPANY_METRICS = [
    ("kk_found", "株式会社の設立", "24-00-16.xlsx", "設立"),
    ("gk_found", "合同会社の設立", "24-00-20.xlsx", "設立"),
    ("kk_dissolve", "株式会社の解散", "24-00-16.xlsx", "解散"),
    ("gk_dissolve", "合同会社の解散", "24-00-20.xlsx", "解散"),
]


def years_inclusive(a: int, b: int) -> list[str]:
    return [str(y) for y in range(a, b + 1)]


def round_n(v: float | None, digits: int) -> float | None:
    if v is None or not isinstance(v, (int, float)) or isinstance(v, bool):
        return None
    f = 10**digits
    return round(v * f) / f


def num(v: object | None) -> float | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and v.strip() in {"…", "...", "－", "-", "—", "―"}:
        return None
    return None


def parse_nengo(text: object | None) -> int | None:
    if not isinstance(text, str):
        return None
    t = re.sub(r"\s+", "", text)
    m = re.match(r"(令和|平成)(元|\d+)年?", t)
    if not m:
        return None
    n = 1 if m.group(2) == "元" else int(m.group(2))
    return (2018 if m.group(1) == "令和" else 1988) + n


class Cube:
    def __init__(self, dims: list[tuple[str, list[str]]], measures: list[str]):
        self.dims = [{"name": n, "codes": codes} for n, codes in dims]
        self.index = [{c: i for i, c in enumerate(codes)} for _, codes in dims]
        self.strides = []
        for i in range(len(dims)):
            s = 1
            for _, codes in dims[i + 1 :]:
                s *= len(codes)
            self.strides.append(s)
        self.size = 1
        for _, codes in dims:
            self.size *= len(codes)
        self.measures = {m: [None] * self.size for m in measures}

    def set(self, measure: str, coords: list[str], value: float | None) -> None:
        offset = 0
        for i, code in enumerate(coords):
            offset += self.index[i][code] * self.strides[i]
        self.measures[measure][offset] = value

    def to_json(self) -> dict:
        return {"dims": self.dims, "measures": self.measures}


def year_cols(ws, header_row: int = 3) -> dict[int, int]:
    out: dict[int, int] = {}
    for c in range(1, ws.max_column + 1):
        year = parse_nengo(ws.cell(header_row, c).value)
        if year is not None:
            out[year] = c
    if not out:
        raise SystemExit(f"{ws.title}: 年次列が読めない")
    return out


def count_row(ws, pred) -> dict[int, float | None]:
    cols = year_cols(ws)
    for i in range(1, ws.max_row + 1):
        lab = ws.cell(i, 2).value
        unit = ws.cell(i, 4).value
        if unit != "件数" or not isinstance(lab, str) or not pred(lab):
            continue
        return {y: num(ws.cell(i, c).value) for y, c in cols.items()}
    raise SystemExit(f"{ws.title}: 件数行が見つからない")


def inherit_row(ws) -> dict[int, float | None]:
    old = count_row(ws, lambda s: s.startswith("相続その他一般承継"))
    new_i = count_row(ws, lambda s: s.startswith("相続による"))
    new_o = count_row(ws, lambda s: s.startswith("一般承継"))
    out: dict[int, float | None] = {}
    for year in set(old) | set(new_i) | set(new_o):
        if year >= 2024:
            a, b = new_i.get(year), new_o.get(year)
            out[year] = None if a is None and b is None else (a or 0) + (b or 0)
        else:
            out[year] = old.get(year)
    return out


def add_series(
    a: dict[int, float | None], b: dict[int, float | None]
) -> dict[int, float | None]:
    out: dict[int, float | None] = {}
    for year in set(a) | set(b):
        x, y = a.get(year), b.get(year)
        out[year] = None if x is None and y is None else (x or 0) + (y or 0)
    return out


def load_totals() -> dict[int, dict[str, float | None]]:
    ws = openpyxl.load_workbook(ESTAT / "24-00-1.xlsx", data_only=True).active
    out: dict[int, dict[str, float | None]] = {}
    era = None
    for row in ws.iter_rows(min_row=6, max_row=ws.max_row, values_only=True):
        era_cell = row[1]
        if era_cell in ("令和", "平成"):
            era = era_cell
        n_cell = row[2]
        if n_cell is None or era is None:
            continue
        s = str(n_cell).strip()
        if s == "元":
            n = 1
        elif s.isdigit():
            n = int(s)
        else:
            continue
        year = (2018 if era == "令和" else 1988) + n
        out[year] = {
            "total": num(row[5]),
            "real_estate": num(row[7]),
            "land": num(row[9]),
            "building": num(row[11]),
        }
    return out


def load_land_building() -> dict[str, dict[int, float | None]]:
    land = openpyxl.load_workbook(ESTAT / "24-00-4.xlsx", data_only=True).active
    bldg = openpyxl.load_workbook(ESTAT / "24-00-5.xlsx", data_only=True).active
    return {
        "sale": add_series(
            count_row(land, lambda s: s.startswith("売買による")),
            count_row(bldg, lambda s: s.startswith("売買による")),
        ),
        "inherit": add_series(inherit_row(land), inherit_row(bldg)),
        "gift": add_series(
            count_row(land, lambda s: s.startswith("遺贈又は贈与")),
            count_row(bldg, lambda s: s.startswith("遺贈又は贈与")),
        ),
        "mortgage": add_series(
            count_row(land, lambda s: s == "抵当権の設定"),
            count_row(bldg, lambda s: s == "抵当権の設定"),
        ),
        "land_sale": count_row(land, lambda s: s.startswith("売買による")),
        "land_inherit": inherit_row(land),
        "land_gift": count_row(land, lambda s: s.startswith("遺贈又は贈与")),
        "land_mortgage": count_row(land, lambda s: s == "抵当権の設定"),
        "land_root": count_row(land, lambda s: s == "根抵当権の設定"),
        "land_preserve": count_row(land, lambda s: s == "所有権の保存"),
        "land_display": count_row(land, lambda s: s.startswith("土地の表示")),
        "land_total": count_row(land, lambda s: s == "総数"),
    }


def load_company(path: str, kind: str) -> dict[int, float | None]:
    ws = openpyxl.load_workbook(ESTAT / path, data_only=True).active
    cols = year_cols(ws)
    for i in range(1, ws.max_row + 1):
        lab = ws.cell(i, 2).value
        unit = ws.cell(i, 4).value
        if lab == kind and unit == "本店":
            return {y: num(ws.cell(i, c).value) for y, c in cols.items()}
    raise SystemExit(f"{path}: {kind}/本店 が見つからない")


def build_era() -> dict:
    years = years_inclusive(ERA_FROM, ERA_TO)
    cube = Cube([("metric", [m[0] for m in METRICS]), ("year", years)], ["counts"])
    totals = load_totals()
    types = load_land_building()
    for y in range(ERA_FROM, ERA_TO + 1):
        ys = str(y)
        t = totals.get(y, {})
        cube.set("counts", ["total", ys], round_n(t.get("total"), 0))
        cube.set("counts", ["real_estate", ys], round_n(t.get("real_estate"), 0))
        cube.set("counts", ["land", ys], round_n(t.get("land"), 0))
        cube.set("counts", ["building", ys], round_n(t.get("building"), 0))
        for code in ("sale", "inherit", "gift", "mortgage"):
            cube.set("counts", [code, ys], round_n(types[code].get(y), 0))
    metrics = [{"code": c, "label": lab, "level": 1, "parent": g} for c, lab, g in METRICS]
    return {**cube.to_json(), "metrics": metrics}


def build_type() -> dict:
    years = years_inclusive(TYPE_FROM, ERA_TO)
    codes = [c for c, _, _ in TYPE_CODES]
    cube = Cube([("code", codes), ("year", years)], ["counts", "share"])
    table = load_land_building()
    key = {
        "sale": "land_sale",
        "inherit": "land_inherit",
        "gift": "land_gift",
        "mortgage": "land_mortgage",
        "root": "land_root",
        "preserve": "land_preserve",
        "display": "land_display",
    }
    for y in range(TYPE_FROM, ERA_TO + 1):
        total = table["land_total"].get(y)
        for code, _label, _pat in TYPE_CODES:
            counts = table[key[code]].get(y)
            share = None if counts is None or total in (None, 0) else counts / total
            cube.set("counts", [code, str(y)], round_n(counts, 0))
            cube.set("share", [code, str(y)], round_n(share, 4))
    return {
        **cube.to_json(),
        "codes": [{"code": c, "label": lab, "level": 1} for c, lab, _ in TYPE_CODES],
    }


def build_company() -> dict:
    years = years_inclusive(TYPE_FROM, ERA_TO)
    cube = Cube(
        [("metric", [m[0] for m in COMPANY_METRICS]), ("year", years)], ["counts"]
    )
    for code, _label, path, kind in COMPANY_METRICS:
        series = load_company(path, kind)
        for y in range(TYPE_FROM, ERA_TO + 1):
            cube.set("counts", [code, str(y)], round_n(series.get(y), 0))
    metrics = [{"code": c, "label": lab, "level": 1} for c, lab, _, _ in COMPANY_METRICS]
    return {**cube.to_json(), "metrics": metrics}


def write_json(name: str, data: dict) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f"{name}.json"
    path.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")))
    print(f"  {name}.json  {path.stat().st_size / 1024:.1f} KB")


def main() -> None:
    print("building cubes")
    write_json("era", build_era())
    write_json("type", build_type())
    write_json("company", build_company())
    print("done")


if __name__ == "__main__":
    main()
